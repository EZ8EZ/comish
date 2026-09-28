"""Turn a Google Sheet (exported as .xlsx, so every tab is included) into row records.

Drive's CSV export only returns the first tab, which would silently drop data. Each
non-empty row becomes a record like "Team: Hawks; Paid: yes; Amount: $100", with the
header row supplying the labels, and is cited as sheet > tab > row N.
"""

import datetime as dt
import io
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class SheetError(ValueError):
    pass


@dataclass(frozen=True)
class SheetRow:
    path: str
    text: str


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_workbook(sheet_name: str, data: bytes) -> list[SheetRow]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise SheetError(f"unreadable spreadsheet: {exc}") from exc
    rows: list[SheetRow] = []
    for tab in workbook.worksheets:
        headers: list[str] | None = None
        for number, raw in enumerate(tab.iter_rows(values_only=True), start=1):
            cells = [_cell(v) for v in raw]
            if not any(cells):
                continue
            if headers is None:
                headers = [c or f"Column {i + 1}" for i, c in enumerate(cells)]
                continue
            parts = [
                f"{headers[i] if i < len(headers) else f'Column {i + 1}'}: {value}"
                for i, value in enumerate(cells)
                if value
            ]
            rows.append(SheetRow(f"{sheet_name} > {tab.title} > row {number}", "; ".join(parts)))
    workbook.close()
    return rows
